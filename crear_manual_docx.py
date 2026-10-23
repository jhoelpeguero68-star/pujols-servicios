# -*- coding: utf-8 -*-
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

doc = docx.Document()

# Margenes
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

def add_heading_1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = RGBColor(10, 37, 64) # #0A2540
    return p

def add_heading_2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(12.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 140, 200)
    return p

def add_p(text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(11)
    r.font.italic = italic
    r.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_bullet(text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(51, 65, 85)
    return p

# Titulo Principal
title_p = doc.add_paragraph()
title_p.paragraph_format.space_after = Pt(2)
title_run = title_p.add_run("MANUAL DE USUARIO - PUJOLS SERVICIOS")
title_run.font.name = 'Calibri'
title_run.font.size = Pt(20)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(10, 37, 64)

sub_p = doc.add_paragraph()
sub_p.paragraph_format.space_after = Pt(16)
sub_run = sub_p.add_run("Sistema de Gestión de Clientes, Citas, Facturación y Equipos\nVersión Oficial en Línea (PWA)")
sub_run.font.name = 'Calibri'
sub_run.font.size = Pt(12)
sub_run.font.italic = True
sub_run.font.color.rgb = RGBColor(100, 116, 139)

# 1. Enlace Oficial
add_heading_1("1. Enlace Oficial Permanente de la Aplicación")
add_p("La aplicación se encuentra publicada y disponible las 24 horas del día en el siguiente enlace oficial permanente:")
link_p = doc.add_paragraph()
link_p.paragraph_format.space_before = Pt(4)
link_p.paragraph_format.space_after = Pt(8)
link_run = link_p.add_run("https://jhoelpeguero68-star.github.io/pujols-servicios/")
link_run.font.name = 'Calibri'
link_run.font.size = Pt(12)
link_run.font.bold = True
link_run.font.color.rgb = RGBColor(0, 102, 204)
add_p("Este enlace es definitivo, no vence, no depende de la computadora y puede abrirse desde cualquier celular (con datos o Wi-Fi) o desde cualquier computadora.")

# 2. Instalacion
add_heading_1("2. Cómo Instalar la App en el Celular (con Logo de la Empresa)")
add_p("PUJOLS Servicios es una aplicación web progresiva (PWA), lo que permite instalarla directamente sin ocupar espacio de descargas pesadas:")
add_heading_2("En teléfonos Android (Google Chrome):")
add_bullet(" Abra el enlace arriba en el navegador Google Chrome.")
add_bullet(" Toque el botón de los 3 puntos verticales en la esquina superior derecha.")
add_bullet(" Seleccione 'Instalar aplicación' o 'Agregar a la pantalla principal'.")
add_bullet(" Confirme en 'Instalar'. En segundos aparecerá el icono oficial con el logo de PUJOLS Servicios en su pantalla de inicio.")

add_heading_2("En teléfonos iPhone (Safari):")
add_bullet(" Abra el enlace en Safari.")
add_bullet(" Toque el botón Compartir (icono cuadrado con una flecha hacia arriba en la barra inferior).")
add_bullet(" Deslice hacia abajo y presione 'Agregar al inicio'.")
add_bullet(" Toque 'Agregar' en la esquina superior derecha.")

# 3. Configuracion Google y WhatsApp
add_heading_1("3. Configuración Inicial (Correo y WhatsApp del Negocio)")
add_p("Para vincular la información de su empresa:")
add_bullet(" En el menú inferior o lateral, ingrese a 'Configuración' (icono de engranaje).")
add_bullet(" En la sección 'Mi empresa', ubique Google Drive o Google Sheets y presione 'Configurar'.")
add_bullet(" Ingrese su correo Gmail para los respaldos en la nube.")
add_bullet(" Escriba su número comercial de WhatsApp (ejemplo: 8095550100) que se usará para la comunicación con los clientes.")
add_bullet(" Opcionalmente pegue los enlaces de su Google Sheets y carpeta de Google Drive.")
add_bullet(" Presione 'Guardar configuración'. Verá los indicadores en verde de vinculación exitosa.")

# 4. Facturación
add_heading_1("4. Cómo Generar e Imprimir Facturas Individuales (PDF / Impresión)")
add_p("La emisión de facturas y recibos de servicio es rápida y genera documentos limpios con el logo de PUJOLS Servicios:")
add_bullet(" Vaya a la pantalla 'Facturación' y presione el botón '+ Generar Factura' (o hágalo desde una orden de trabajo finalizada).")
add_bullet(" Seleccione el cliente y detalle el concepto del servicio prestado (ej. Instalación Split Inverter 18,000 BTU, vacío y soporte).")
add_bullet(" Ingrese el monto acordado (puede colocar cualquier valor exacto sin restricciones) y seleccione si está Pagada o Pendiente.")
add_bullet(" Presione 'Generar y Previsualizar' para ver la factura con el logo oficial.")
add_bullet(" Para guardar o imprimir: Presione 'Imprimir / PDF'. La hoja se imprimirá de manera individual y limpia, sin barras de menú ni elementos de fondo.")
add_bullet(" Para enviar al cliente: Presione 'Compartir por WhatsApp'. Se abrirá el mensaje redactado para el cliente con el monto y número de factura.")

# 5. WhatsApp con Mensajes de Cortesía
add_heading_1("5. Prueba y Envío de Mensajes por WhatsApp (con Cortesía)")
add_p("Todos los mensajes automáticos inician con el saludo de cortesía formal correspondiente:")
add_bullet(" Para clientes particulares: Inicia con 'Estimado/a Sr./Sra. [Nombre del Cliente], reciba un cordial saludo de PUJOLS Servicios...'.")
add_bullet(" Para empresas o constructoras: Inicia con 'Estimados señores de [Nombre], reciban un atento saludo de PUJOLS Servicios...'.")
add_p("Para probar el envío con su propio teléfono:")
add_bullet(" 1. Vaya a 'Clientes' y registre su propio número de teléfono real.")
add_bullet(" 2. En 'Servicios', pulse 'WhatsApp' en una cita para ver el mensaje automático con fecha, hora, equipo y técnico.")
add_bullet(" 3. En 'Recordatorios', pulse 'WhatsApp' en un aire acondicionado para enviar el aviso preventivo de los 7 u 8 meses.")
add_bullet(" 4. En 'Facturación', pulse 'WhatsApp' para enviar el detalle de pago.")

# 6. Datos Demo vs Reales
add_heading_1("6. Datos de Demostración y Cómo Empezar con Clientes Reales")
add_bullet(" La app incluye clientes y servicios de prueba para que pueda practicar y conocer todas las funciones antes de usarla.")
add_bullet(" Cuando decida cargar sus datos reales: Vaya a 'Configuración' -> en la parte inferior pulse el botón rojo 'Limpiar datos de demostración (Empezar de cero)'.")
add_bullet(" La aplicación quedará completamente en blanco y lista para registrar sus clientes reales.")
add_bullet(" Si más adelante desea volver a ver datos de ejemplo para demostración, puede presionar 'Restaurar datos de ejemplo'.")

# 7. Modo sin Internet
add_heading_1("7. Funcionamiento sin Internet (Modo Offline)")
add_p("Si se encuentra en un sótano o lugar sin señal:")
add_bullet(" Puede seguir abriendo la aplicación, registrando nuevos clientes, tomando ubicación GPS y creando órdenes.")
add_bullet(" Los datos se guardan de forma segura en la memoria de su teléfono o PC.")
add_bullet(" En cuanto su dispositivo recupere conexión Wi-Fi o datos móviles, la app sincronizará automáticamente los registros pendientes.")

out_path = r"c:\Users\Admin\Desktop\MANUAL_PUJOLS_SERVICIOS_ACTUALIZADO.docx"
doc.save(out_path)
print("DOCX guardado exitosamente en:", out_path)
